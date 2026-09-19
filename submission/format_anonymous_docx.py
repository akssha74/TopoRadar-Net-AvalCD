#!/usr/bin/env python3
"""Apply explicit JMS-compatible page and typography settings to the DOCX."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_run_font(run, size=10.0, bold=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_fixed_layout(table):
    table.autofit = False
    table_pr = table._tbl.tblPr
    layout = table_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_column_widths(table, widths_mm):
    """Set both the OOXML grid and cell widths so renderers honor the layout."""
    set_fixed_layout(table)
    grid_columns = table._tbl.tblGrid.gridCol_lst
    for grid_column, width in zip(grid_columns, widths_mm):
        grid_column.set(qn("w:w"), str(int(Mm(width).twips)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_mm):
            cell.width = Mm(width)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "path",
        nargs="?",
        default="JMS_26-11479_Main_Manuscript_Anonymous.docx",
    )
    args = parser.parse_args()
    path = Path(args.path)
    document = Document(path)

    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(25.4)
        section.bottom_margin = Mm(25.4)
        section.left_margin = Mm(25.4)
        section.right_margin = Mm(25.4)
        section.header_distance = Mm(12.7)
        section.footer_distance = Mm(12.7)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        if style_name not in styles:
            continue
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(10)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.0

    in_references = False
    for paragraph in document.paragraphs:
        is_reference_heading = paragraph.text.strip().lower() == "references"
        if is_reference_heading:
            in_references = True
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = (
            0.95 if in_references and not is_reference_heading else 1.0
        )
        run_size = 9.0 if in_references and not is_reference_heading else 10.0
        for run in paragraph.runs:
            set_run_font(run, run_size)

    header_rows = (1, 2, 1, 2, 2, 2, 1, 1)
    widths_mm = {
        1: (40, 12, 18, 18, 18, 18, 17, 17),  # Table 2
        4: (50, 18, 18, 18, 18, 18, 18),      # Table 5
        5: (15, 45, 14, 17, 17, 17, 17, 17),  # Table 6
    }
    for table_index, table in enumerate(document.tables):
        if table_index in widths_mm:
            set_column_widths(table, widths_mm[table_index])
        else:
            table.autofit = True
        for row_index, row in enumerate(table.rows):
            prevent_row_split(row)
            if row_index < header_rows[table_index]:
                set_repeat_table_header(row)
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        set_run_font(
                            run,
                            8.0 if table_index in widths_mm else 8.5,
                        )

    document.save(path)
    print(
        f"Formatted {path}: A4, 25.4 mm margins, Times New Roman 10 pt, "
        "single-spaced, 0 pt paragraph-after; tables 8–8.5 pt with repeated "
        "headers and non-splitting rows."
    )


if __name__ == "__main__":
    main()
