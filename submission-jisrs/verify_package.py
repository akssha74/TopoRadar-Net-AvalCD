#!/usr/bin/env python3
"""Run the mechanical JISRS package and scientific reproduction gates."""

from __future__ import annotations

import json
import re
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SUBMISSION = ROOT / "submission-jisrs"
PAPER = ROOT / "paper-jisrs"


def run(command: list[str], cwd: Path) -> str:
    result = subprocess.run(
        command,
        cwd=cwd,
        text=True,
        check=True,
        capture_output=True,
    )
    return result.stdout + result.stderr


def main() -> None:
    print(run([sys.executable, "check_compliance.py"], SUBMISSION))
    print(run([sys.executable, "experiments/code/reproduce_all_metrics.py"], ROOT))
    print(run([sys.executable, "reviews/harness.py"], ROOT))

    metadata = json.loads((SUBMISSION / "submission_metadata.json").read_text())
    required = [
        ROOT / metadata["main_document"],
        ROOT / metadata["title_page"],
        ROOT / metadata["supplementary_information"],
        ROOT / metadata["cover_letter"],
        ROOT / metadata["highlights"],
        ROOT / metadata["source_package"],
        PAPER / "main.pdf",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise SystemExit(f"Missing submission files: {missing}")

    anonymous_text = (PAPER / "main.tex").read_text(encoding="utf-8")
    anonymous_text += (PAPER / "supplementary_information.tex").read_text(
        encoding="utf-8"
    )
    with zipfile.ZipFile(SUBMISSION / "JISRS_Main_Manuscript_Anonymous.docx") as docx:
        document_xml_bytes = docx.read("word/document.xml")
        anonymous_text += document_xml_bytes.decode("utf-8")
        docx_document_xml = document_xml_bytes.decode("utf-8")
        core = docx.read("docProps/core.xml").decode("utf-8")
        footer_xml = "\n".join(
            docx.read(name).decode("utf-8")
            for name in docx.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )

    identity_terms = (
        "Akshay",
        "Lalji",
        "SAGE University",
        "akssha74",
        "@gmail",
    )
    leaks = [term for term in identity_terms if term.lower() in anonymous_text.lower()]
    if leaks:
        raise SystemExit(f"Anonymous-package identity leaks: {leaks}")
    if re.search(r"<dc:creator>[^<]+", core):
        raise SystemExit("Anonymous DOCX has a non-empty creator")
    for required_text in (
        "Abstract",
        "Cloud cover",
        "Keywords",
        "Sentinel-1 SAR",
        "snow avalanche mapping",
    ):
        if required_text not in docx_document_xml:
            raise SystemExit(f"Anonymous DOCX missing required text: {required_text}")

    namespaces = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    }
    root = ET.fromstring(document_xml_bytes)
    paragraph_records: list[tuple[str, bytes]] = []
    text_tags = {f"{{{namespaces['w']}}}t", f"{{{namespaces['m']}}}t"}
    for paragraph in root.iter(f"{{{namespaces['w']}}}p"):
        text = "".join(
            node.text or "" for node in paragraph.iter() if node.tag in text_tags
        )
        paragraph_records.append((text, ET.tostring(paragraph)))

    table2_caption = next(
        (text for text, _ in paragraph_records if text.startswith("Table 2.")),
        None,
    )
    if table2_caption is None or "±" not in table2_caption or "30%" not in table2_caption:
        raise SystemExit("DOCX Table 2 caption lost its ± or 30% math content")

    affine_record = next(
        (
            (text, xml)
            for text, xml in paragraph_records
            if text.startswith("The affine valid-mask area is")
        ),
        None,
    )
    if affine_record is None:
        raise SystemExit("DOCX lacks the affine-area paragraph")
    affine_text, affine_xml = affine_record
    if "km²" not in affine_text or "km⁻²" not in affine_text:
        raise SystemExit("DOCX affine-area paragraph lacks contiguous Unicode units")
    affine_element = ET.fromstring(affine_xml)
    if any(
        node.tag == f"{{{namespaces['w']}}}i" for node in affine_element.iter()
    ):
        raise SystemExit("DOCX affine-area unit paragraph contains italic runs")

    document = Document(SUBMISSION / "JISRS_Main_Manuscript_Anonymous.docx")
    if not document.sections:
        raise SystemExit("DOCX has no section")
    section = document.sections[0]
    if round(section.page_width.mm) != 210 or round(section.page_height.mm) != 297:
        raise SystemExit("DOCX is not A4")
    for section_index, docx_section in enumerate(document.sections, start=1):
        line_numbers = docx_section._sectPr.find(qn("w:lnNumType"))
        if line_numbers is None:
            raise SystemExit(
                f"DOCX section {section_index} lacks automatic line numbering"
            )
        if (
            line_numbers.get(qn("w:countBy")) != "1"
            or line_numbers.get(qn("w:start")) != "1"
            or line_numbers.get(qn("w:restart")) != "continuous"
        ):
            raise SystemExit(
                f"DOCX section {section_index} line numbering is not continuous"
            )
    if not re.search(r"<w:instrText[^>]*>\s*PAGE\s*</w:instrText>", footer_xml):
        raise SystemExit("DOCX lacks an automatic PAGE field in its footer")
    if abs(document.styles["Normal"].font.size.pt - 12) > 0.01:
        raise SystemExit("DOCX Normal style is not 12 pt")
    docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for caption_prefix in ("Table 1.", "Table 2.", "Table 3.", "Fig. 1 ", "Fig. 2 "):
        if caption_prefix not in docx_text:
            raise SystemExit(f"DOCX missing caption prefix: {caption_prefix}")
    if "Fig. 1." in docx_text or "Fig. 2." in docx_text:
        raise SystemExit("DOCX figure caption contains punctuation after its number")
    if "Statements and Declarations" in docx_text:
        raise SystemExit(
            "Blinded DOCX still contains the prohibited Statements and Declarations section"
        )
    for prohibited_declaration in (
        "Funding. No external grant funding",
        "Competing interests.",
        "Ethics approval and consent.",
        "Consent for publication.",
        "Author contributions.",
    ):
        if prohibited_declaration in docx_text:
            raise SystemExit(
                f"Blinded DOCX retains declaration text: {prohibited_declaration}"
            )

    title_page = Document(SUBMISSION / "JISRS_Title_Page.docx")
    title_page_text = "\n".join(
        paragraph.text for paragraph in title_page.paragraphs
    )
    if "Declarations" not in title_page_text:
        raise SystemExit("Title page lacks a Declarations heading")
    if "Highlights" in title_page_text:
        raise SystemExit("Title page duplicates the separately supplied Highlights")
    required_title_headings = [
        "Conflict of Interest",
        "Funding",
        "Author’s Contribution",
        "Acknowledgement",
    ]
    title_heading_positions = [
        title_page_text.find(heading) for heading in required_title_headings
    ]
    if any(position < 0 for position in title_heading_positions):
        raise SystemExit(
            "Title page lacks one or more editorially required declaration headings"
        )
    if title_heading_positions != sorted(title_heading_positions):
        raise SystemExit(
            "Title page declaration headings are not in the editorially required order"
        )
    if (
        "The authors declare that they have no conflict of interest."
        not in title_page_text
    ):
        raise SystemExit("Title page lacks the required Conflict of Interest statement")
    if (
        "No funding was received for conducting this study or preparing this manuscript."
        not in title_page_text
    ):
        raise SystemExit("Title page lacks the required Funding statement")
    acknowledgement_text = title_page_text.split("Acknowledgement", 1)[1].split(
        "Ethics Approval and Consent", 1
    )[0]
    if "Not applicable." not in acknowledgement_text:
        raise SystemExit("Title page Acknowledgement is not marked Not applicable")
    if Path(metadata["supplementary_information"]).name != "ESM_1.pdf":
        raise SystemExit("Supplementary Information is not named ESM_1.pdf")

    with tempfile.TemporaryDirectory() as temp:
        result = subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                temp,
                str(SUBMISSION / "JISRS_Main_Manuscript_Anonymous.docx"),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        rendered = Path(temp) / "JISRS_Main_Manuscript_Anonymous.pdf"
        if not rendered.exists():
            raise SystemExit(f"LibreOffice did not render DOCX: {result.stdout}")
        extracted = subprocess.run(
            ["pdftotext", str(rendered), "-"],
            check=True,
            capture_output=True,
            text=True,
        ).stdout
        if "❑" in extracted or "□" in extracted:
            raise SystemExit("Rendered DOCX contains missing-glyph boxes")
        normalized = re.sub(r"\s+", " ", extracted)
        for malformed_unit in ("k m2", "k m 2", "h a k m", "ha k m"):
            if malformed_unit in normalized:
                raise SystemExit(
                    f"Rendered DOCX contains separated unit letters: {malformed_unit}"
                )
        if not all(
            unit in extracted for unit in ("km²", "km⁻²")
        ):
            raise SystemExit("Rendered DOCX lacks contiguous superscript unit text")
        if any(line.lstrip().startswith(",") for line in extracted.splitlines()):
            raise SystemExit("Rendered DOCX contains a line-start orphan comma")
        page_info = subprocess.run(
            ["pdfinfo", str(rendered)],
            check=True,
            capture_output=True,
            text=True,
        ).stdout
        pages_match = re.search(r"(?m)^Pages:\s+(\d+)", page_info)
        if not pages_match or int(pages_match.group(1)) not in (12, 13):
            raise SystemExit("Rendered DOCX does not satisfy the 12-13 page target")

        title_result = subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                temp,
                str(SUBMISSION / "JISRS_Title_Page.docx"),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        rendered_title = Path(temp) / "JISRS_Title_Page.pdf"
        if not rendered_title.exists():
            raise SystemExit(
                f"LibreOffice did not render title page: {title_result.stdout}"
            )
        title_page_info = subprocess.run(
            ["pdfinfo", str(rendered_title)],
            check=True,
            capture_output=True,
            text=True,
        ).stdout
        title_pages_match = re.search(r"(?m)^Pages:\s+(\d+)", title_page_info)
        if not title_pages_match or int(title_pages_match.group(1)) != 1:
            raise SystemExit("Identified title-page DOCX does not render to one page")

        stored_preview = SUBMISSION / "JISRS_Main_Manuscript_Anonymous.pdf"
        stored_text = subprocess.run(
            ["pdftotext", str(stored_preview), "-"],
            check=True,
            capture_output=True,
            text=True,
        ).stdout
        if re.sub(r"\s+", " ", stored_text).strip() != re.sub(
            r"\s+", " ", extracted
        ).strip():
            raise SystemExit("Stored DOCX preview text differs from a fresh render")

        fresh_prefix = Path(temp) / "fresh"
        stored_prefix = Path(temp) / "stored"
        subprocess.run(
            ["pdftoppm", "-png", "-r", "96", str(rendered), str(fresh_prefix)],
            check=True,
            capture_output=True,
        )
        subprocess.run(
            [
                "pdftoppm",
                "-png",
                "-r",
                "96",
                str(stored_preview),
                str(stored_prefix),
            ],
            check=True,
            capture_output=True,
        )
        fresh_pages = sorted(Path(temp).glob("fresh-*.png"))
        stored_pages = sorted(Path(temp).glob("stored-*.png"))
        if len(fresh_pages) != len(stored_pages) or any(
            fresh.read_bytes() != stored.read_bytes()
            for fresh, stored in zip(fresh_pages, stored_pages)
        ):
            raise SystemExit("Stored DOCX preview raster differs from a fresh render")

    with zipfile.ZipFile(SUBMISSION / "source_package.zip") as source:
        names = set(source.namelist())
        required_members = {
            "TopoRadar-Net-JISRS/main.tex",
            "TopoRadar-Net-JISRS/supplementary_information.tex",
            "TopoRadar-Net-JISRS/references.bib",
            "TopoRadar-Net-JISRS/sn-jnl.cls",
            "TopoRadar-Net-JISRS/sn-apacite.bst",
            "TopoRadar-Net-JISRS/figures/Fig1.eps",
            "TopoRadar-Net-JISRS/figures/Fig2.eps",
            "TopoRadar-Net-JISRS/evidence/event_sensor_manifest.json",
        }
        absent = sorted(required_members - names)
        if absent:
            raise SystemExit(f"Source package missing members: {absent}")

    print("PASS: JISRS package, anonymity, Word, source, and scientific gates")


if __name__ == "__main__":
    main()
