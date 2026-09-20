#!/usr/bin/env python3
"""Apply JISRS double-blind Word formatting and structural controls."""

from __future__ import annotations

import argparse
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]


def extract_braced(text: str, command: str) -> str:
    marker = f"\\{command}"
    start = text.index(marker) + len(marker)
    while text[start].isspace():
        start += 1
    depth = 0
    for index in range(start, len(text)):
        if text[index] == "{" and text[index - 1] != "\\":
            depth += 1
        elif text[index] == "}" and text[index - 1] != "\\":
            depth -= 1
            if depth == 0:
                return text[start + 1 : index]
    raise ValueError(f"Unterminated {command}")


def latex_to_plain(text: str) -> str:
    result = subprocess.run(
        ["pandoc", "-f", "latex", "-t", "plain"],
        input=text,
        text=True,
        check=True,
        capture_output=True,
    )
    return " ".join(result.stdout.split())


def insert_before(anchor: Paragraph, text: str, style: str) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def prefix_caption(paragraph: Paragraph, prefix: str) -> None:
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_properties.append(OxmlElement("w:b"))
    run.append(run_properties)
    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = prefix
    run.append(text)
    insert_at = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insert_at, run)


def replace_interoperability_sensitive_paragraphs(document: Document) -> None:
    replacements = {
        "The full architecture is not the best pooled configuration": (
            "The full architecture is not the best pooled configuration (Table 3). "
            "Removing aspect and its fixed-reference coordinate or removing "
            "cross-attention increases pooled macro F1 to 65.35% and 65.07%, "
            "respectively, mainly through gains in Pish. In Tromsø, seed-42 paired "
            "block tests favour the full model over No-LIA by 2.04 percentage points "
            "(p<0.001) and No-CrossAttn by 1.42 percentage points (p=0.010). These "
            "comparisons establish component sensitivity, not a causal "
            "terrain-correction mechanism."
        ),
        "The affine valid-mask area is": (
            "The affine valid-mask area is 196.820 km² in Tromsø and 267.435 km² "
            "in Pish. In Tromsø, TopoRadar-Net produces 74.77 ha of false positives "
            "at 0.38 ha km⁻². In Pish, it produces 129.69 ha at 0.48 ha km⁻². "
            "Relative to the self-attention U-Net, this is a 30.0% reduction in Pish "
            "(129.69 versus 185.32 ha) but a 28.5% increase in Tromsø (74.77 versus "
            "58.19 ha). Validation-calibrated road-corridor alerts achieve 83.0% "
            "precision and 62.5% recall in Pish. The single road-intersecting "
            "reference avalanche in Tromsø is missed by every TopoRadar-Net seed. "
            "All cached complete-scene runs are below 5 s on Apple MPS hardware, "
            "but timing excludes upstream acquisition and human review."
        ),
    }
    for paragraph in document.paragraphs:
        for start, replacement in replacements.items():
            if paragraph.text.strip().startswith(start):
                paragraph.clear()
                paragraph.add_run(replacement)
                break


def set_run_font(run, size: float) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Times New Roman"
    )
    run.font.size = Pt(size)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def prevent_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_continuous_line_numbers(section) -> None:
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:lnNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:start"), "1")
    line_numbers.set(qn("w:restart"), "continuous")
    sect_pr.append(line_numbers)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    args = parser.parse_args()

    document = Document(args.path)
    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(20)
        section.right_margin = Mm(20)
        add_continuous_line_numbers(section)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Times New Roman"
    )
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        if style_name not in document.styles:
            continue
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Times New Roman"
        )
        style.font.size = Pt(12)
        style.font.bold = True
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(0)

    if not any(p.text.strip() == "Abstract" for p in document.paragraphs):
        source = (ROOT / "paper-jisrs" / "main.tex").read_text(encoding="utf-8")
        abstract = latex_to_plain(extract_braced(source, "abstract"))
        keywords = ", ".join(
            item.strip() for item in extract_braced(source, "keywords").split(",")
        )
        introduction = next(
            p for p in document.paragraphs if p.text.strip() == "Introduction"
        )
        insert_before(introduction, "Abstract", "Heading 1")
        insert_before(introduction, abstract, "Normal")
        insert_before(introduction, "Keywords", "Heading 1")
        insert_before(introduction, keywords, "Normal")

    caption_prefixes = {
        "AvalCD scenes used": "Table 1. ",
        "TopoRadar-Net schematic": "Fig. 1. ",
        "Zero-shot test performance": "Table 2. ",
        "Zero-shot pixel F1": "Fig. 2. ",
        "Ablation performance": "Table 3. ",
    }
    for paragraph in document.paragraphs:
        stripped = paragraph.text.strip()
        for start, prefix in caption_prefixes.items():
            if stripped.startswith(start) and not stripped.startswith(prefix):
                prefix_caption(paragraph, prefix)
                break

    replace_interoperability_sensitive_paragraphs(document)

    for paragraph in document.paragraphs:
        paragraph.paragraph_format.line_spacing = 2.0
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_run_font(run, 12)

    for table in document.tables:
        table.autofit = True
        for row_index, row in enumerate(table.rows):
            prevent_split(row)
            if row_index == 0:
                set_repeat_header(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        set_run_font(run, 9)

    printable_width = Mm(170)
    for shape in document.inline_shapes:
        ratio = shape.height / shape.width
        shape.width = printable_width
        shape.height = int(printable_width * ratio)

    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    document.save(args.path)

    print(
        f"Formatted {args.path}: A4, 20-mm margins, 12-pt Times New Roman, "
        "double-spaced, continuous line numbers, non-splitting table rows."
    )


if __name__ == "__main__":
    main()
